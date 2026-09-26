"""Unit tests for journal parse, posting, layout, artifacts, and export fail-closed."""

from __future__ import annotations

import base64
import io

import pytest
from openpyxl import Workbook

from app.services import artifact_store
from app.services.ingest import IngestError, decode_file_base64, resolve_xlsx_bytes
from app.services.journal_parse import ParseError, parse_nhat_ky_chung_bytes
from app.services.layout import workbook_matrices
from app.services.sheets_export import ExportError, export_t_account_workbook
from app.services.t_accounts import PostError, build_t_accounts


@pytest.fixture(autouse=True)
def artifact_tmpdir(tmp_path, monkeypatch):
    monkeypatch.setenv("ACCOUNTING_ARTIFACT_DIR", str(tmp_path / "arts"))
    return tmp_path


def _sample_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Nhật ký chung"
    ws["A1"] = "CÔNG TY TNHH DEMO – MST: 0123456789"
    ws["A3"] = "SỔ NHẬT KÝ CHUNG"
    ws["A4"] = "Từ 2023-01-01 đến 2023-12-31"
    headers = [
        "Mã dòng",
        "Phân hệ",
        "Ngày hạch toán",
        "Ngày chứng từ",
        "Mã chứng từ",
        "Diễn giải",
        "TK nợ",
        "Đối tượng nợ",
        "TK có",
        "Đối tượng có",
        "Giá trị",
        "Tags",
    ]
    for i, h in enumerate(headers, start=1):
        ws.cell(5, i, h)
    rows = [
        ("2023-01-05", "CT1", "Mua hàng nhập kho", "156", "331", 1000),
        ("2023-01-06", "CT2", "Xuất kho giá vốn", "632", "156", 400),
        ("2023-01-07", "CT3", "Mua NVL", "152", "331", 250),
        ("2023-01-08", "CT4", "Xuất NVL SX", "154", "152", 100),
    ]
    for r, (d, ct, desc, no, co, amt) in enumerate(rows, start=6):
        ws.cell(r, 3, d)
        ws.cell(r, 5, ct)
        ws.cell(r, 6, desc)
        ws.cell(r, 7, no)
        ws.cell(r, 9, co)
        ws.cell(r, 11, amt)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_and_balance():
    parsed = parse_nhat_ky_chung_bytes(_sample_xlsx())
    assert parsed["entryCount"] == 4
    assert parsed["mst"] == "0123456789"
    built = build_t_accounts(parsed["entries"], circular="TT200")
    assert built["balanced"] is True
    assert built["totalDebit"] == built["totalCredit"] == 1750.0
    by = {a["account"]: a for a in built["accounts"]}
    assert by["156"]["debitTotal"] == 1000
    assert by["156"]["creditTotal"] == 400
    assert by["156"]["net"] == 600


def test_books_unbalanced_fails():
    bad = [
        {
            "debit_account": "156",
            "credit_account": "331",
            "debit_amount": 100,
            "credit_amount": 40,
            "amount": 100,
        }
    ]
    with pytest.raises(PostError) as ei:
        build_t_accounts(bad)
    assert ei.value.code == "BOOKS_UNBALANCED"


def test_one_sided_line_unbalanced():
    bad = [{"debit_account": "156", "credit_account": "", "amount": 100}]
    with pytest.raises(PostError) as ei:
        build_t_accounts(bad)
    assert ei.value.code == "BOOKS_UNBALANCED"


def test_no_posted_entries_fails():
    with pytest.raises(PostError) as ei:
        build_t_accounts(
            [
                {"debit_account": "", "credit_account": "", "amount": 0},
                {"debit_account": "x", "credit_account": "y", "amount": -1},
            ]
        )
    assert ei.value.code == "NO_POSTED_ENTRIES"


def test_header_missing():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "hello"
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ParseError) as ei:
        parse_nhat_ky_chung_bytes(buf.getvalue())
    assert ei.value.code == "HEADER_NOT_FOUND"


def test_layout_tabs():
    parsed = parse_nhat_ky_chung_bytes(_sample_xlsx())
    built = build_t_accounts(parsed["entries"])
    mats = workbook_matrices(built)
    assert "T-accounts" in mats and "Sơ đồ chữ T" in mats
    assert any("TK 156" in str(r) for r in mats["T-accounts"])
    assert mats["Sơ đồ chữ T"][0]


def test_ingest_base64_and_resolve():
    b64 = base64.b64encode(_sample_xlsx()).decode("ascii")
    raw = decode_file_base64(b64)
    assert raw[:2] == b"PK"
    data, src = resolve_xlsx_bytes(file_base64=b64)
    assert src == "file_base64"
    with pytest.raises(IngestError):
        resolve_xlsx_bytes(file_base64=b64, attachment_id="x")


def test_artifact_tenant_isolation():
    parsed = parse_nhat_ky_chung_bytes(_sample_xlsx())
    aid = artifact_store.new_id("nkc")
    artifact_store.save_json(aid, {"kind": "nhat_ky_chung", **parsed}, tenant_id="tenant-a")
    loaded = artifact_store.load_json(aid, tenant_id="tenant-a")
    assert loaded["entryCount"] == 4
    with pytest.raises(artifact_store.ArtifactError) as ei:
        artifact_store.load_json(aid, tenant_id="tenant-b")
    assert ei.value.code == "ARTIFACT_NOT_FOUND"


def test_export_credentials_required_when_not_dry_run():
    parsed = parse_nhat_ky_chung_bytes(_sample_xlsx())
    built = build_t_accounts(parsed["entries"])
    with pytest.raises(ExportError) as ei:
        export_t_account_workbook(built, dry_run=False)
    assert ei.value.code == "CREDENTIALS_REQUIRED"
    ok = export_t_account_workbook(built, dry_run=True)
    assert ok["dryRun"] is True
    assert "T-accounts" in ok["matrices"]

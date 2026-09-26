"""Parse Vietnamese Sổ nhật ký chung .xlsx into journal entries."""

from __future__ import annotations

import io
import re
from dataclasses import asdict, dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Optional

from openpyxl import load_workbook

REQUIRED_HEADERS = {
    "tk_no": ("tk nợ", "tk no", "debit account", "tài khoản nợ"),
    "tk_co": ("tk có", "tk co", "credit account", "tài khoản có"),
    "gia_tri": ("giá trị", "gia tri", "amount", "số tiền", "so tien"),
}

OPTIONAL_HEADERS = {
    "ngay": ("ngày hạch toán", "ngay hach toan", "ngày chứng từ", "date"),
    "dien_giai": ("diễn giải", "dien giai", "description", "narration"),
    "ma_ct": ("mã chứng từ", "ma chung tu", "voucher", "doc no"),
}


class ParseError(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass
class JournalEntry:
    line: int
    date: Optional[str]
    debit_account: str
    credit_account: str
    amount: float
    description: str
    voucher: str


def _norm_header(v: Any) -> str:
    s = str(v or "").strip().lower()
    s = s.replace("\n", " ")
    s = re.sub(r"\s+", " ", s)
    return s


def _cell_str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()


def _cell_amount(v: Any) -> Optional[float]:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace(" ", "")
    if not s:
        return None
    try:
        return float(Decimal(s))
    except (InvalidOperation, ValueError):
        return None


def _match_col(header_row: list[Any], aliases: tuple[str, ...]) -> Optional[int]:
    norms = [_norm_header(h) for h in header_row]
    for i, h in enumerate(norms):
        for a in aliases:
            if h == a or a in h:
                return i
    return None


def _find_header_row(rows: list[tuple[Any, ...]], max_scan: int = 30) -> tuple[int, dict[str, int]]:
    for idx, row in enumerate(rows[:max_scan]):
        mapping: dict[str, int] = {}
        for key, aliases in REQUIRED_HEADERS.items():
            col = _match_col(list(row), aliases)
            if col is None:
                mapping = {}
                break
            mapping[key] = col
        if not mapping:
            continue
        for key, aliases in OPTIONAL_HEADERS.items():
            col = _match_col(list(row), aliases)
            if col is not None:
                mapping[key] = col
        return idx, mapping
    raise ParseError(
        "HEADER_NOT_FOUND",
        "Could not find header row with TK nợ / TK có / Giá trị (or English aliases)",
    )


def _normalize_account(raw: str) -> str:
    s = (raw or "").strip()
    # keep leading account code digits + optional suffix
    m = re.match(r"^([0-9]{2,4}[A-Za-z0-9]*)", s)
    if m:
        return m.group(1)
    return s


def parse_nhat_ky_chung_bytes(data: bytes, sheet_name: Optional[str] = None) -> dict[str, Any]:
    if not data:
        raise ParseError("EMPTY_FILE", "xlsx bytes are empty")
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ParseError("INVALID_XLSX", f"Cannot open workbook: {exc}") from exc

    ws = None
    if sheet_name:
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            for name in wb.sheetnames:
                if "nhật" in name.lower() or "nhat" in name.lower() or "ký chung" in name.lower():
                    ws = wb[name]
                    break
    if ws is None:
        for name in wb.sheetnames:
            if "nhật" in name.lower() or "nhat ky" in name.lower().replace("ý", "y"):
                ws = wb[name]
                break
    if ws is None:
        ws = wb[wb.sheetnames[0]]

    sheet_title = ws.title
    rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    if not rows:
        raise ParseError("EMPTY_SHEET", "worksheet has no rows")

    # Company / MST from title rows
    mst = ""
    company = ""
    period = ""
    for row in rows[:6]:
        text = " ".join(_cell_str(c) for c in row if c is not None)
        if not company and ("công ty" in text.lower() or "cong ty" in text.lower()):
            company = text[:200]
        m = re.search(r"MST[:\s]*([0-9]{8,14})", text, re.I)
        if m:
            mst = m.group(1)
        if "từ" in text.lower() or "from" in text.lower() or "đến" in text.lower():
            period = text[:200]

    header_idx, cols = _find_header_row(rows)
    entries: list[JournalEntry] = []
    skipped = 0
    for i, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        def at(key: str) -> Any:
            c = cols.get(key)
            if c is None or c >= len(row):
                return None
            return row[c]

        debit = _normalize_account(_cell_str(at("tk_no")))
        credit = _normalize_account(_cell_str(at("tk_co")))
        amount = _cell_amount(at("gia_tri"))
        if not debit or not credit or amount is None:
            skipped += 1
            continue
        if amount == 0:
            skipped += 1
            continue
        entries.append(
            JournalEntry(
                line=i,
                date=_cell_str(at("ngay")) or None,
                debit_account=debit,
                credit_account=credit,
                amount=abs(amount),
                description=_cell_str(at("dien_giai")),
                voucher=_cell_str(at("ma_ct")),
            )
        )

    if not entries:
        raise ParseError("NO_ENTRIES", "No journal lines with TK nợ / TK có / Giá trị found")

    dates = [e.date for e in entries if e.date]
    sample = [asdict(e) for e in entries[:5]]
    return {
        "sheetName": sheet_title,
        "company": company,
        "mst": mst,
        "periodHint": period,
        "entryCount": len(entries),
        "skippedRows": skipped,
        "dateMin": min(dates) if dates else None,
        "dateMax": max(dates) if dates else None,
        "sampleEntries": sample,
        "entries": [asdict(e) for e in entries],
    }
